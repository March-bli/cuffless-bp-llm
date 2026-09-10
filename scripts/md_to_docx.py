"""将 thesis_final.md 转换为 Word 文档（.docx）"""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "/home/lby/projects/bishe/docs/thesis_final.md"
DST = "/home/lby/projects/bishe/docs/thesis_final.docx"

doc = Document()

# 设置默认字体
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

def add_runs_with_bold(paragraph, text):
    """解析 **bold** 标记并添加 runs。"""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # 去掉行内公式标记
            part = part.replace("$", "")
            paragraph.add_run(part)

def add_table(lines):
    """解析 markdown 表格并添加到 doc。"""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            # 跳过分隔行
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            rows.append(cells)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            text = cell.replace("**", "")
            table.cell(i, j).text = text

lines = open(SRC, encoding="utf-8").read().splitlines()

i = 0
in_table = False
table_lines = []

while i < len(lines):
    line = lines[i]

    # 空行
    if not line.strip():
        in_table = False
        i += 1
        continue

    # 表格
    if line.strip().startswith("|"):
        in_table = True
        table_lines.append(line)
        # 收集连续的表格行
        while i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            i += 1
            table_lines.append(lines[i])
        add_table(table_lines)
        table_lines = []
        i += 1
        continue

    in_table = False

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)", line)
    if m:
        level = len(m.group(1))
        text = m.group(2).replace("**", "")
        doc.add_heading(text, level=level)
        i += 1
        continue

    # 列表
    if line.strip().startswith("- "):
        text = line.strip()[2:]
        p = doc.add_paragraph(style="List Bullet")
        add_runs_with_bold(p, text)
        i += 1
        continue
    m = re.match(r"^\s*(\d+)[\.\)]\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs_with_bold(p, m.group(2))
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    add_runs_with_bold(p, line)
    i += 1

doc.save(DST)
print(f"✅ Word 文档已生成: {DST}")
