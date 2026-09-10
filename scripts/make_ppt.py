"""生成毕设技术汇报 PPT（英文，给导师看）"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

TITLE_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x1F, 0x77, 0xB4)
RED = RGBColor(0xD6, 0x27, 0x28)

FIG = "/home/lby/projects/bishe/docs/figures"


def add_slide(title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    if title:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
        tf = box.text_frame
        tf.text = title
        p = tf.paragraphs[0]
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = TITLE_COLOR
    return slide


def add_bullets(slide, items, left=0.6, top=1.4, width=12.0, height=5.5, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            p.text = item[0]
            p.level = item[1]
            p.font.bold = item[2]
            if len(item) > 3:
                p.font.color.rgb = item[3]
        else:
            p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(8)
    return box


def add_table(slide, rows, left=0.6, top=1.5, width=12.0, font_size=16):
    n_rows, n_cols = len(rows), len(rows[0])
    table = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                   Inches(width), Inches(0.5 * n_rows)).table
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = str(cell)
            table.cell(i, j).text_frame.paragraphs[0].font.size = Pt(font_size)
            if i == 0:
                table.cell(i, j).text_frame.paragraphs[0].font.bold = True
    return table


# ── Slide 1: 标题 ───────────────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(2.5))
tf = box.text_frame
tf.text = "Novel Blood Pressure Monitoring Using\nWearable Devices and Artificial Intelligence"
for i, p in enumerate(tf.paragraphs):
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR
box2 = slide.shapes.add_textbox(Inches(1), Inches(5.0), Inches(11.3), Inches(1))
tf2 = box2.text_frame
tf2.text = "Boyuan Li  |  Supervisor: Dr. Shaoxiong Sun\nProject Progress Report"
for p in tf2.paragraphs:
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)

# ── Slide 2: 问题背景 ──────────────────────────────────────────────────
s = add_slide("Background: The Cross-Subject Generalisation Challenge")
add_bullets(s, [
    "Goal: cuffless BP monitoring from wrist-worn PPG sensors",
    "Core problem: PPG-to-BP mapping varies across individuals",
    ("→ Traditional ML fails under subject-level splits", 1, True, RED),
    "",
    "XGBoost under different evaluation protocols:",
    "• Segment-level split (data leakage): SBP MAE = 8.0 mmHg  ← misleading",
    "• Subject-level split (correct):      SBP MAE = 17.0 mmHg  ← real difficulty",
])

# ── Slide 3: 数据集 ────────────────────────────────────────────────────
s = add_slide("Dataset: PulseDB (MIMIC + Vital)")
add_table(s, [
    ["Item", "Value"],
    ["Subjects", "857"],
    ["PPG segments", "148,012"],
    ["Signal", "PPG / ECG / ABP @ 125 Hz"],
    ["Segment length", "10 seconds"],
    ["Features", "55-dimensional + Age/Gender"],
    ["Split", "Strict subject-level (20 test subjects)"],
], top=1.6)

# ── Slide 4: 方法一 ────────────────────────────────────────────────────
s = add_slide("Method 1: Cross-Subject Zero-Shot")
add_bullets(s, [
    "LLM receives (features, BP, demographics) from OTHER subjects",
    "Then predicts BP for a NEW, unseen subject",
    "",
    "Result: SBP 20.8 mmHg, DBP 12.6 mmHg",
    ("≈ XGBoost (17.0/11.6) → individual physiology is the real bottleneck", 1, True, ACCENT),
])

# ── Slide 5: 方法二（核心） + K 定义 ───────────────────────────────────
s = add_slide("Method 2: Few-Shot Personalised Calibration (Core)")
add_bullets(s, [
    ("What is K?", 0, True, RED),
    "K = number of calibration samples from the TARGET subject",
    "Each sample = (55 PPG features, true SBP/DBP) pair",
    "Randomly selected from the subject's own segments (not 'first K' chronologically)",
    "",
    "LLM sees K calibration samples + 1 new sample → predicts BP",
    "Ablation: K ∈ {1, 3, 5, 10, 20}",
])

# ── Slide 6: 核心结果曲线 ──────────────────────────────────────────────
s = add_slide("Core Result: Calibration Curve")
s.shapes.add_picture(f"{FIG}/fig1_calibration_curve.png", Inches(1.5), Inches(1.4), width=Inches(10.3))

# ── Slide 7: 消融表格 ─────────────────────────────────────────────────
s = add_slide("Ablation Results")
add_table(s, [
    ["K", "SBP MAE", "DBP MAE", "Note"],
    ["0 (zero-shot)", "20.8", "12.6", "cross-subject baseline"],
    ["1", "21.8", "18.8", "1 sample insufficient"],
    ["3", "14.4", "11.6", "improving"],
    ["5", "9.3", "6.1", "BHS Grade B"],
    ["10", "8.3", "4.1", "★ DBP Grade A"],
    ["20", "8.4", "4.5", "saturated"],
], top=1.7, font_size=18)

# ── Slide 8: 与传统 ML 对比 ────────────────────────────────────────────
s = add_slide("Comparison: vs Traditional ML")
s.shapes.add_picture(f"{FIG}/fig2_comparison.png", Inches(2.2), Inches(1.4), width=Inches(8.9))

# ── Slide 9: 与 LLM 对比 ──────────────────────────────────────────────
s = add_slide("Key Finding: Calibration Samples > Model Scale")
add_table(s, [
    ["Method", "Model scale", "SBP MAE", "DBP MAE"],
    ["LLM-BP (2024)", "GPT-4 API", "9.3", "6.4"],
    ["DeepSeek API (K=5)", "671B", "8.9", "4.3"],
    ["Ours (K=10)", "8B local", "8.3", "4.1"],
], top=2.0, font_size=18)
add_bullets(s, [
    ("8B local model with K=10 matches/beats 671B model", 0, True, RED),
    "→ Personal calibration data outweighs model capacity",
], top=4.2, size=18)

# ── Slide 10: 语义抽象 ────────────────────────────────────────────────
s = add_slide("Semantic Abstraction (SensorLM / TimeSRL)")
s.shapes.add_picture(f"{FIG}/fig3_semantic.png", Inches(2.2), Inches(1.4), width=Inches(8.9))
add_bullets(s, [
    ("Rule-based semantic descriptions improve DBP (5.6 vs 6.1)", 0, True, ACCENT),
    ("Pure qualitative LLM summaries collapse (26.9) → must retain quantitative anchors", 0, True, RED),
], top=5.8, size=16)

# ── Slide 11: 关键发现总结 ────────────────────────────────────────────
s = add_slide("Summary of Findings")
add_bullets(s, [
    "1. Cross-subject zero-shot ≈ XGBoost → individual physiology is the bottleneck",
    "2. Few-shot calibration (K=10): SBP 8.3 / DBP 4.1 — DBP meets BHS Grade A",
    "3. 51% / 65% improvement over XGBoost",
    "4. Calibration sample size matters more than model scale (8B ≈ 671B)",
    "5. Semantic descriptions help DBP; over-abstraction collapses",
])

# ── Slide 12: 下一步 ───────────────────────────────────────────────────
s = add_slide("Next Steps")
add_bullets(s, [
    "• Writing dissertation (all experiments complete)",
    "• Preparing presentation for 4 September",
    "• Feedback welcome on the K=10 personalisation story",
])

prs.save("/home/lby/projects/bishe/docs/project_progress_slides.pptx")
print("✅ PPT 已生成: project_progress_slides.pptx")
